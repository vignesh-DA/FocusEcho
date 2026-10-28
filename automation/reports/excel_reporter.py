"""
Excel Report Generator for FocusEcho E2E Selenium Framework
Generates enterprise multi-sheet styled workbooks using openpyxl:
1. Automation_Test_Report.xlsx (6 sheets: Executed, Passed, Failed, Skipped, Metrics, Defects)
2. Failed_Test_Cases.xlsx
3. Passed_Test_Cases.xlsx
4. Summary_Report.xlsx
"""

from typing import List
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from automation.models.test_result import TestCaseResult
from automation.config.config import EXCEL_DIR
from automation.utils.logger import logger

class ExcelReporter:
    """Creates styled Excel reports conforming strictly to QA specifications."""

    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    ACCENT_FILL = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")

    PASS_FILL = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    PASS_FONT = Font(name="Calibri", size=10, color="065F46", bold=True)

    FAIL_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    FAIL_FONT = Font(name="Calibri", size=10, color="991B1B", bold=True)

    SKIP_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    SKIP_FONT = Font(name="Calibri", size=10, color="92400E", bold=True)

    THIN_BORDER = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB")
    )

    @classmethod
    def generate_all(cls, results: List[TestCaseResult], base_url: str) -> None:
        """Generate all 4 required Excel reports."""
        cls.generate_master_report(results, base_url)
        cls.generate_passed_report(results)
        cls.generate_failed_report(results)
        cls.generate_summary_report(results, base_url)
        logger.info(f"Generated 4 Excel workbooks in {EXCEL_DIR}")

    @classmethod
    def _apply_row_styling(cls, ws, row_idx: int, status: str) -> None:
        for cell in ws[row_idx]:
            cell.border = cls.THIN_BORDER
            cell.alignment = Alignment(vertical="center")
            if cell.value in ("Passed", "Failed", "Skipped", "Blocked"):
                if cell.value == "Passed":
                    cell.fill = cls.PASS_FILL
                    cell.font = cls.PASS_FONT
                elif cell.value == "Failed":
                    cell.fill = cls.FAIL_FILL
                    cell.font = cls.FAIL_FONT
                else:
                    cell.fill = cls.SKIP_FILL
                    cell.font = cls.SKIP_FONT

    @classmethod
    def _auto_fit_columns(cls, ws) -> None:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    @classmethod
    def generate_master_report(cls, results: List[TestCaseResult], base_url: str) -> None:
        """Automation_Test_Report.xlsx with 6 sheets."""
        wb = openpyxl.Workbook()
        
        # Sheet 1: Executed Test Cases
        ws1 = wb.active
        ws1.title = "Executed Test Cases"
        headers1 = ["Test ID", "Module", "Test Name", "Status", "Execution Time (s)", "Priority"]
        ws1.append(headers1)
        for c in ws1[1]:
            c.font = cls.HEADER_FONT
            c.fill = cls.HEADER_FILL

        for r in results:
            ws1.append([r.test_id, r.module, r.test_name, r.status, r.execution_time, r.priority])
            cls._apply_row_styling(ws1, ws1.max_row, r.status)
        cls._auto_fit_columns(ws1)

        # Sheet 2: Passed Tests
        ws2 = wb.create_sheet(title="Passed Tests")
        headers2 = ["Test ID", "Module", "Test Name", "Execution Time (s)", "Priority"]
        ws2.append(headers2)
        for c in ws2[1]:
            c.font = cls.HEADER_FONT
            c.fill = cls.ACCENT_FILL
        for r in results:
            if r.status == "Passed":
                ws2.append([r.test_id, r.module, r.test_name, r.execution_time, r.priority])
                cls._apply_row_styling(ws2, ws2.max_row, "Passed")
        cls._auto_fit_columns(ws2)

        # Sheet 3: Failed Tests
        ws3 = wb.create_sheet(title="Failed Tests")
        headers3 = ["Test ID", "Module", "Test Name", "Priority", "Error Message", "Stack Trace"]
        ws3.append(headers3)
        for c in ws3[1]:
            c.font = cls.HEADER_FONT
            c.fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
        for r in results:
            if r.status == "Failed":
                ws3.append([r.test_id, r.module, r.test_name, r.priority, r.error_message or "", r.stack_trace or ""])
                cls._apply_row_styling(ws3, ws3.max_row, "Failed")
        cls._auto_fit_columns(ws3)

        # Sheet 4: Skipped Tests
        ws4 = wb.create_sheet(title="Skipped Tests")
        headers4 = ["Test ID", "Module", "Test Name", "Reason"]
        ws4.append(headers4)
        for c in ws4[1]:
            c.font = cls.HEADER_FONT
            c.fill = cls.HEADER_FILL
        for r in results:
            if r.status in ("Skipped", "Blocked"):
                ws4.append([r.test_id, r.module, r.test_name, r.error_message or "Skipped"])
                cls._apply_row_styling(ws4, ws4.max_row, r.status)
        cls._auto_fit_columns(ws4)

        # Sheet 5: Execution Metrics
        ws5 = wb.create_sheet(title="Execution Metrics")
        total = len(results)
        passed = sum(1 for r in results if r.status == "Passed")
        failed = sum(1 for r in results if r.status == "Failed")
        skipped = sum(1 for r in results if r.status in ("Skipped", "Blocked"))
        pass_pct = (passed / total * 100) if total else 0
        total_time = sum(r.execution_time for r in results)

        ws5.append(["Metric", "Value"])
        ws5.append(["Live Deployment URL", base_url])
        ws5.append(["Execution Timestamp", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")])
        ws5.append(["Total Test Cases", total])
        ws5.append(["Passed Test Cases", passed])
        ws5.append(["Failed Test Cases", failed])
        ws5.append(["Skipped Test Cases", skipped])
        ws5.append(["Pass Percentage", f"{pass_pct:.2f}%"])
        ws5.append(["Total Execution Time (s)", f"{total_time:.2f}"])
        ws5.append(["Average Execution Time per Test", f"{(total_time / total):.3f}s" if total else "0s"])
        for c in ws5[1]:
            c.font = cls.HEADER_FONT
            c.fill = cls.HEADER_FILL
        cls._auto_fit_columns(ws5)

        # Sheet 6: Defect Summary
        ws6 = wb.create_sheet(title="Defect Summary")
        ws6.append(["Defect ID", "Associated Test ID", "Module", "Severity", "Summary", "Status"])
        for c in ws6[1]:
            c.font = cls.HEADER_FONT
            c.fill = cls.HEADER_FILL
        failed_tests = [r for r in results if r.status == "Failed"]
        if not failed_tests:
            ws6.append(["-", "None", "All Modules", "None", "Zero defects encountered in this test run.", "Closed"])
        else:
            for idx, r in enumerate(failed_tests, 1):
                ws6.append([f"DEF-{idx:03d}", r.test_id, r.module, r.priority, r.error_message or "Assertion Error", "Open"])
        cls._auto_fit_columns(ws6)

        target_file = EXCEL_DIR / "Automation_Test_Report.xlsx"
        wb.save(target_file)

    @classmethod
    def generate_passed_report(cls, results: List[TestCaseResult]) -> None:
        """Passed_Test_Cases.xlsx."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Passed Test Cases"
        ws.append(["Test ID", "Module", "Test Name", "Execution Time (s)", "Priority", "Status"])
        for c in ws[1]:
            c.font = cls.HEADER_FONT
            c.fill = cls.ACCENT_FILL

        for r in results:
            if r.status == "Passed":
                ws.append([r.test_id, r.module, r.test_name, r.execution_time, r.priority, "Passed"])
                cls._apply_row_styling(ws, ws.max_row, "Passed")
        cls._auto_fit_columns(ws)
        wb.save(EXCEL_DIR / "Passed_Test_Cases.xlsx")

    @classmethod
    def generate_failed_report(cls, results: List[TestCaseResult]) -> None:
        """Failed_Test_Cases.xlsx."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Failed Test Cases"
        ws.append(["Test ID", "Module", "Test Name", "Priority", "Status", "Failure Reason"])
        for c in ws[1]:
            c.font = cls.HEADER_FONT
            c.fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")

        failed = [r for r in results if r.status == "Failed"]
        if not failed:
            ws.append(["None", "-", "No test failures", "-", "-", "100% Passing"])
        else:
            for r in failed:
                ws.append([r.test_id, r.module, r.test_name, r.priority, "Failed", r.error_message or ""])
                cls._apply_row_styling(ws, ws.max_row, "Failed")
        cls._auto_fit_columns(ws)
        wb.save(EXCEL_DIR / "Failed_Test_Cases.xlsx")

    @classmethod
    def generate_summary_report(cls, results: List[TestCaseResult], base_url: str) -> None:
        """Summary_Report.xlsx."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Executive Summary"
        ws.append(["Module", "Total Tests", "Passed", "Failed", "Pass Rate (%)", "Avg Duration (s)"])
        for c in ws[1]:
            c.font = cls.HEADER_FONT
            c.fill = cls.HEADER_FILL

        modules = sorted(list(set(r.module for r in results)))
        for mod in modules:
            mod_tests = [r for r in results if r.module == mod]
            tot = len(mod_tests)
            pas = sum(1 for r in mod_tests if r.status == "Passed")
            fai = sum(1 for r in mod_tests if r.status == "Failed")
            pct = (pas / tot * 100) if tot else 0
            avg_dur = sum(r.execution_time for r in mod_tests) / tot if tot else 0
            ws.append([mod, tot, pas, fai, f"{pct:.1f}%", f"{avg_dur:.3f}"])
        cls._auto_fit_columns(ws)
        wb.save(EXCEL_DIR / "Summary_Report.xlsx")
